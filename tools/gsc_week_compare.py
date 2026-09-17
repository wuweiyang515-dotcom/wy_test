#!/usr/bin/env python3
"""GSC 周对比分析：找出流量升降最多的 URL，并归因到关键词的四个维度。

指标口径
--------
* 展现量 / 点击量：按天求和。
* 点击率：clicks / impressions（加权，不是日均 CTR 的算术平均）。
* 平均排名：SUM(rank * impressions) / SUM(impressions)（按展现加权）。
* URL 级指标来自 url_gsc_traffic 中的 url_* 字段。由于同一 (url_id, data_date)
  会在每个关键词行上重复一次 URL 级数值，取数时先按 (url_id, data_date) 去重。
* 关键词级指标来自 keyword_* 字段，仅覆盖 GSC 返回的头部查询，
  合计会小于 URL 级总量（匿名化查询不下发），因此只用于“归因”，不用于“总量”。

变化归因（乘法分解）
--------------------
clicks = impressions * ctr，于是
    Δclicks = Δimpressions * ctr_prev      (曝光效应)
            + impressions_prev * Δctr      (点击率效应)
            + Δimpressions * Δctr          (交互项)
点击率效应再结合 Δ平均排名判断是不是排名变化导致的。

凭证只从环境变量读取：DB_HOST / DB_PORT / DB_DATABASE / DB_USERNAME / DB_PASSWORD。
"""

from __future__ import annotations

import argparse
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RATE_SCALE = 100.0  # CTR 以百分比展示
RANK_MOVE_THRESHOLD = 0.5  # 平均排名变化超过该值才视为发生了排名迁移


# --------------------------------------------------------------------------- #
# 基础工具
# --------------------------------------------------------------------------- #
def parse_date(value: str) -> date:
    return datetime.strptime(value, "%Y-%m-%d").date()


def safe_div(numerator: float, denominator: float) -> float:
    return numerator / denominator if denominator else 0.0


def decompose(impr_prev: float, clicks_prev: float, impr_cur: float, clicks_cur: float) -> dict:
    """把 Δclicks 拆成曝光效应 / 点击率效应 / 交互项。"""
    ctr_prev = safe_div(clicks_prev, impr_prev)
    ctr_cur = safe_div(clicks_cur, impr_cur)
    d_impr = impr_cur - impr_prev
    d_ctr = ctr_cur - ctr_prev
    return {
        "ctr_prev": ctr_prev,
        "ctr_cur": ctr_cur,
        "d_impr": d_impr,
        "d_ctr": d_ctr,
        "d_clicks": clicks_cur - clicks_prev,
        "impr_effect": d_impr * ctr_prev,
        "ctr_effect": impr_prev * d_ctr,
        "interaction": d_impr * d_ctr,
    }


def main_driver(parts: dict) -> str:
    """判断主要驱动因素。"""
    candidates = {
        "展现量": parts["impr_effect"],
        "点击率": parts["ctr_effect"],
        "曝光×点击率交互": parts["interaction"],
    }
    name, value = max(candidates.items(), key=lambda kv: abs(kv[1]))
    if abs(value) < 1e-9:
        return "无明显变化"
    return name


def rank_comment(d_rank: float) -> str:
    """平均排名变化说明（GSC position 越小越好）。"""
    if d_rank <= -RANK_MOVE_THRESHOLD:
        return f"排名前进 {abs(d_rank):.1f} 位"
    if d_rank >= RANK_MOVE_THRESHOLD:
        return f"排名后退 {d_rank:.1f} 位"
    return "排名基本持平"


def driver_sentence(parts: dict, d_rank: float) -> str:
    driver = main_driver(parts)
    if driver == "无明显变化":
        return "四个维度均无明显变化"
    if driver == "点击率":
        if abs(d_rank) >= RANK_MOVE_THRESHOLD:
            return f"点击率变化主导，且{rank_comment(d_rank)}，属于排名驱动"
        return "点击率变化主导，但排名基本持平，更可能是标题/摘要/SERP 形态或搜索意图变化"
    if driver == "展现量":
        return f"展现量变化主导（需求或收录/覆盖面变化），{rank_comment(d_rank)}"
    return f"展现量与点击率同向共振，{rank_comment(d_rank)}"


# --------------------------------------------------------------------------- #
# 取数
# --------------------------------------------------------------------------- #
def connect():
    password = os.environ.get("DB_PASSWORD")
    if not password:
        sys.exit("缺少环境变量 DB_PASSWORD（其余：DB_HOST/DB_PORT/DB_DATABASE/DB_USERNAME）")
    return pymysql.connect(
        host=os.environ.get("DB_HOST", "127.0.0.1"),
        port=int(os.environ.get("DB_PORT", "3306")),
        user=os.environ.get("DB_USERNAME", "root"),
        password=password,
        database=os.environ.get("DB_DATABASE", "pmauto_system"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=30,
        read_timeout=900,
    )


def resolve_domain_id(conn, domain_like: str) -> int:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT id, url FROM domains WHERE url LIKE %s ORDER BY id LIMIT 2",
            (f"%{domain_like}%",),
        )
        rows = cur.fetchall()
    if not rows:
        sys.exit(f"找不到域名：{domain_like}")
    if len(rows) > 1:
        sys.exit(f"域名匹配到多条，请用 --domain-id 指定：{[r['url'] for r in rows]}")
    return int(rows[0]["id"])


def fetch_url_level(conn, domain_id, prefixes, prev_start, cur_start, cur_end):
    """URL 级两周指标。先按 (url_id, data_date) 去重再聚合。"""
    like_sql = " OR ".join(["u.url LIKE %s"] * len(prefixes))
    sql = f"""
        SELECT x.url_id,
               u.url                                       AS url,
               SUM(IF(x.period = 1, x.impressions, 0))     AS cur_impr,
               SUM(IF(x.period = 1, x.clicks, 0))          AS cur_clicks,
               SUM(IF(x.period = 1, x.rank_weighted, 0))   AS cur_rank_w,
               SUM(IF(x.period = 0, x.impressions, 0))     AS prev_impr,
               SUM(IF(x.period = 0, x.clicks, 0))          AS prev_clicks,
               SUM(IF(x.period = 0, x.rank_weighted, 0))   AS prev_rank_w
        FROM (
            SELECT t.url_id,
                   t.data_date,
                   IF(t.data_date >= %s, 1, 0)               AS period,
                   MAX(t.url_show_num)                       AS impressions,
                   MAX(t.url_click_num)                      AS clicks,
                   MAX(t.url_avg_rank) * MAX(t.url_show_num) AS rank_weighted
            FROM url_gsc_traffic t
            WHERE t.domain_id = %s
              AND t.data_date BETWEEN %s AND %s
            GROUP BY t.url_id, t.data_date
        ) x
        JOIN urls u ON u.id = x.url_id
        WHERE ({like_sql})
        GROUP BY x.url_id, u.url
        HAVING cur_impr > 0 OR prev_impr > 0
    """
    params = [cur_start, domain_id, prev_start, cur_end] + [f"{p}%" for p in prefixes]
    with conn.cursor() as cur:
        cur.execute(sql, params)
        return cur.fetchall()


def fetch_keyword_level(conn, domain_id, url_ids, prev_start, cur_start, cur_end):
    """指定 URL 的关键词级两周指标。"""
    if not url_ids:
        return []
    placeholders = ",".join(["%s"] * len(url_ids))
    sql = f"""
        SELECT t.url_id,
               t.keyword_md5,
               MAX(t.keyword)                                                        AS keyword,
               SUM(IF(t.data_date >= %s, t.keyword_show_num, 0))                     AS cur_impr,
               SUM(IF(t.data_date >= %s, t.keyword_click_num, 0))                    AS cur_clicks,
               SUM(IF(t.data_date >= %s, t.keyword_avg_rank * t.keyword_show_num, 0)) AS cur_rank_w,
               SUM(IF(t.data_date <  %s, t.keyword_show_num, 0))                     AS prev_impr,
               SUM(IF(t.data_date <  %s, t.keyword_click_num, 0))                    AS prev_clicks,
               SUM(IF(t.data_date <  %s, t.keyword_avg_rank * t.keyword_show_num, 0)) AS prev_rank_w
        FROM url_gsc_traffic t
        WHERE t.domain_id = %s
          AND t.data_date BETWEEN %s AND %s
          AND t.keyword_md5 <> ''
          AND t.url_id IN ({placeholders})
        GROUP BY t.url_id, t.keyword_md5
        HAVING cur_impr > 0 OR prev_impr > 0
    """
    params = [cur_start] * 6 + [domain_id, prev_start, cur_end] + list(url_ids)
    with conn.cursor() as cur:
        cur.execute(sql, params)
        return cur.fetchall()


# --------------------------------------------------------------------------- #
# 计算
# --------------------------------------------------------------------------- #
def group_of(url: str, prefixes: list[str]) -> str:
    """最长前缀归组，避免重复计数。"""
    best = ""
    for prefix in prefixes:
        if url.startswith(prefix) and len(prefix) > len(best):
            best = prefix
    return best or "(其他)"


def build_url_rows(raw_rows, prefixes):
    rows = []
    for r in raw_rows:
        cur_impr = float(r["cur_impr"])
        prev_impr = float(r["prev_impr"])
        parts = decompose(prev_impr, float(r["prev_clicks"]), cur_impr, float(r["cur_clicks"]))
        cur_rank = safe_div(float(r["cur_rank_w"]), cur_impr)
        prev_rank = safe_div(float(r["prev_rank_w"]), prev_impr)
        d_rank = cur_rank - prev_rank if (cur_impr and prev_impr) else 0.0
        rows.append(
            {
                "url_id": r["url_id"],
                "url": r["url"],
                "group": group_of(r["url"], prefixes),
                "prev_impr": prev_impr,
                "cur_impr": cur_impr,
                "prev_clicks": float(r["prev_clicks"]),
                "cur_clicks": float(r["cur_clicks"]),
                "prev_ctr": parts["ctr_prev"],
                "cur_ctr": parts["ctr_cur"],
                "prev_rank": prev_rank,
                "cur_rank": cur_rank,
                "d_impr": parts["d_impr"],
                "d_clicks": parts["d_clicks"],
                "d_ctr": parts["d_ctr"],
                "d_rank": d_rank,
                "impr_effect": parts["impr_effect"],
                "ctr_effect": parts["ctr_effect"],
                "interaction": parts["interaction"],
                "driver": main_driver(parts),
                "explain": driver_sentence(parts, d_rank),
            }
        )
    return rows


def build_keyword_rows(raw_rows, url_lookup):
    rows = []
    for r in raw_rows:
        cur_impr = float(r["cur_impr"])
        prev_impr = float(r["prev_impr"])
        parts = decompose(prev_impr, float(r["prev_clicks"]), cur_impr, float(r["cur_clicks"]))
        cur_rank = safe_div(float(r["cur_rank_w"]), cur_impr)
        prev_rank = safe_div(float(r["prev_rank_w"]), prev_impr)
        if prev_impr == 0:
            status = "新增词"
        elif cur_impr == 0:
            status = "流失词"
        else:
            status = "留存词"
        d_rank = cur_rank - prev_rank if status == "留存词" else 0.0
        rows.append(
            {
                "url": url_lookup.get(r["url_id"], str(r["url_id"])),
                "url_id": r["url_id"],
                "keyword": (r["keyword"] or "").strip(),
                "status": status,
                "prev_impr": prev_impr,
                "cur_impr": cur_impr,
                "prev_clicks": float(r["prev_clicks"]),
                "cur_clicks": float(r["cur_clicks"]),
                "prev_ctr": parts["ctr_prev"],
                "cur_ctr": parts["ctr_cur"],
                "prev_rank": prev_rank,
                "cur_rank": cur_rank,
                "d_impr": parts["d_impr"],
                "d_clicks": parts["d_clicks"],
                "d_ctr": parts["d_ctr"],
                "d_rank": d_rank,
                "impr_effect": parts["impr_effect"],
                "ctr_effect": parts["ctr_effect"],
                "interaction": parts["interaction"],
                "driver": main_driver(parts) if status == "留存词" else status,
            }
        )
    return rows


def aggregate(rows, key_func):
    buckets = defaultdict(
        lambda: {
            "urls": 0,
            "prev_impr": 0.0,
            "cur_impr": 0.0,
            "prev_clicks": 0.0,
            "cur_clicks": 0.0,
            "prev_rank_w": 0.0,
            "cur_rank_w": 0.0,
        }
    )
    for row in rows:
        b = buckets[key_func(row)]
        b["urls"] += 1
        b["prev_impr"] += row["prev_impr"]
        b["cur_impr"] += row["cur_impr"]
        b["prev_clicks"] += row["prev_clicks"]
        b["cur_clicks"] += row["cur_clicks"]
        b["prev_rank_w"] += row["prev_rank"] * row["prev_impr"]
        b["cur_rank_w"] += row["cur_rank"] * row["cur_impr"]
    out = []
    for key, b in buckets.items():
        parts = decompose(b["prev_impr"], b["prev_clicks"], b["cur_impr"], b["cur_clicks"])
        prev_rank = safe_div(b["prev_rank_w"], b["prev_impr"])
        cur_rank = safe_div(b["cur_rank_w"], b["cur_impr"])
        out.append(
            {
                "key": key,
                "urls": b["urls"],
                "prev_impr": b["prev_impr"],
                "cur_impr": b["cur_impr"],
                "prev_clicks": b["prev_clicks"],
                "cur_clicks": b["cur_clicks"],
                "prev_ctr": parts["ctr_prev"],
                "cur_ctr": parts["ctr_cur"],
                "prev_rank": prev_rank,
                "cur_rank": cur_rank,
                "d_impr": parts["d_impr"],
                "d_clicks": parts["d_clicks"],
                "d_ctr": parts["d_ctr"],
                "d_rank": cur_rank - prev_rank,
                "impr_effect": parts["impr_effect"],
                "ctr_effect": parts["ctr_effect"],
                "interaction": parts["interaction"],
                "driver": main_driver(parts),
            }
        )
    return sorted(out, key=lambda r: -r["cur_clicks"])


# --------------------------------------------------------------------------- #
# 输出
# --------------------------------------------------------------------------- #
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def write_sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = (widths or {}).get(
            header, max(12, min(46, len(header) * 2 + 6))
        )
    return ws


URL_HEADERS = [
    "URL",
    "分组",
    "上周展现",
    "本周展现",
    "展现差值",
    "上周点击",
    "本周点击",
    "点击差值",
    "上周点击率%",
    "本周点击率%",
    "点击率差值(pp)",
    "上周平均排名",
    "本周平均排名",
    "排名差值",
    "展现效应(点击)",
    "点击率效应(点击)",
    "交互项(点击)",
    "主因",
    "结论",
]


def url_row_values(r):
    return [
        r["url"],
        r["group"],
        int(r["prev_impr"]),
        int(r["cur_impr"]),
        int(r["d_impr"]),
        int(r["prev_clicks"]),
        int(r["cur_clicks"]),
        int(r["d_clicks"]),
        round(r["prev_ctr"] * RATE_SCALE, 2),
        round(r["cur_ctr"] * RATE_SCALE, 2),
        round(r["d_ctr"] * RATE_SCALE, 2),
        round(r["prev_rank"], 2),
        round(r["cur_rank"], 2),
        round(r["d_rank"], 2),
        round(r["impr_effect"], 1),
        round(r["ctr_effect"], 1),
        round(r["interaction"], 1),
        r["driver"],
        r["explain"],
    ]


KW_HEADERS = [
    "URL",
    "关键词",
    "词状态",
    "上周展现",
    "本周展现",
    "展现差值",
    "上周点击",
    "本周点击",
    "点击差值",
    "上周点击率%",
    "本周点击率%",
    "点击率差值(pp)",
    "上周平均排名",
    "本周平均排名",
    "排名差值",
    "展现效应(点击)",
    "点击率效应(点击)",
    "交互项(点击)",
    "主因",
]


def kw_row_values(r):
    return [
        r["url"],
        r["keyword"],
        r["status"],
        int(r["prev_impr"]),
        int(r["cur_impr"]),
        int(r["d_impr"]),
        int(r["prev_clicks"]),
        int(r["cur_clicks"]),
        int(r["d_clicks"]),
        round(r["prev_ctr"] * RATE_SCALE, 2),
        round(r["cur_ctr"] * RATE_SCALE, 2),
        round(r["d_ctr"] * RATE_SCALE, 2),
        round(r["prev_rank"], 2),
        round(r["cur_rank"], 2),
        round(r["d_rank"], 2),
        round(r["impr_effect"], 1),
        round(r["ctr_effect"], 1),
        round(r["interaction"], 1),
        r["driver"],
    ]


SUMMARY_HEADERS = [
    "范围",
    "URL 数",
    "上周展现",
    "本周展现",
    "展现差值",
    "上周点击",
    "本周点击",
    "点击差值",
    "上周点击率%",
    "本周点击率%",
    "点击率差值(pp)",
    "上周平均排名",
    "本周平均排名",
    "排名差值",
    "展现效应(点击)",
    "点击率效应(点击)",
    "交互项(点击)",
    "主因",
]


def summary_row_values(r):
    return [
        r["key"],
        r["urls"],
        int(r["prev_impr"]),
        int(r["cur_impr"]),
        int(r["d_impr"]),
        int(r["prev_clicks"]),
        int(r["cur_clicks"]),
        int(r["d_clicks"]),
        round(r["prev_ctr"] * RATE_SCALE, 2),
        round(r["cur_ctr"] * RATE_SCALE, 2),
        round(r["d_ctr"] * RATE_SCALE, 2),
        round(r["prev_rank"], 2),
        round(r["cur_rank"], 2),
        round(r["d_rank"], 2),
        round(r["impr_effect"], 1),
        round(r["ctr_effect"], 1),
        round(r["interaction"], 1),
        r["driver"],
    ]


# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(description="GSC 周对比与关键词归因分析")
    parser.add_argument("--domain", default="eufy.com", help="域名关键字，用于在 domains 表中定位")
    parser.add_argument("--domain-id", type=int, help="直接指定 domain_id，优先于 --domain")
    parser.add_argument(
        "--paths",
        nargs="+",
        default=[
            "https://www.eufy.com/eu-fr/",
            "https://www.eufy.com/eu-fr/blogs/",
            "https://www.eufy.com/eu-fr/collections/",
        ],
        help="要分析的 URL 前缀（按最长前缀归组）",
    )
    parser.add_argument("--week-start", default="2026-08-15", help="本周开始日期 YYYY-MM-DD")
    parser.add_argument("--week-end", default="2026-08-21", help="本周结束日期 YYYY-MM-DD")
    parser.add_argument("--top", type=int, default=20, help="升/降榜各取多少条 URL")
    parser.add_argument("--min-impr", type=int, default=100, help="进入主榜的最低展现量（两周取较大值）")
    parser.add_argument("--kw-top", type=int, default=15, help="每个 URL 展开多少个关键词")
    parser.add_argument("--out", default="gsc_week_compare.xlsx", help="输出 Excel 路径")
    parser.add_argument("--markdown", help="可选：把文字结论写入该 Markdown 文件")
    args = parser.parse_args()

    cur_start = parse_date(args.week_start)
    cur_end = parse_date(args.week_end)
    if cur_end < cur_start:
        sys.exit("--week-end 不能早于 --week-start")
    span = (cur_end - cur_start).days + 1
    prev_end = cur_start - timedelta(days=1)
    prev_start = prev_end - timedelta(days=span - 1)

    conn = connect()
    try:
        domain_id = args.domain_id or resolve_domain_id(conn, args.domain)
        url_rows = build_url_rows(
            fetch_url_level(conn, domain_id, args.paths, prev_start, cur_start, cur_end),
            args.paths,
        )

        qualified = [r for r in url_rows if max(r["prev_impr"], r["cur_impr"]) >= args.min_impr]
        gainers = [r for r in sorted(qualified, key=lambda r: -r["d_clicks"])[: args.top] if r["d_clicks"] > 0]
        losers = [r for r in sorted(qualified, key=lambda r: r["d_clicks"])[: args.top] if r["d_clicks"] < 0]
        impr_gainers = sorted(qualified, key=lambda r: -r["d_impr"])[: args.top]
        impr_losers = sorted(qualified, key=lambda r: r["d_impr"])[: args.top]

        url_lookup = {r["url_id"]: r["url"] for r in url_rows}
        focus_ids = [r["url_id"] for r in gainers + losers]
        kw_rows = build_keyword_rows(
            fetch_keyword_level(conn, domain_id, focus_ids, prev_start, cur_start, cur_end),
            url_lookup,
        )
    finally:
        conn.close()

    kw_by_url = defaultdict(list)
    for row in kw_rows:
        kw_by_url[row["url_id"]].append(row)

    def kw_slice(url_list, ascending: bool):
        out = []
        for u in url_list:
            words = sorted(
                kw_by_url.get(u["url_id"], []),
                key=lambda r: r["d_clicks"] if ascending else -r["d_clicks"],
            )
            picked = [w for w in words if (w["d_clicks"] < 0 if ascending else w["d_clicks"] > 0)]
            if not picked:  # 点击无变化时退回按展现排序，仍能解释原因
                picked = sorted(words, key=lambda r: r["d_impr"] if ascending else -r["d_impr"])
            out.extend(picked[: args.kw_top])
        return out

    # ---------------- Excel ----------------
    wb = Workbook()
    wb.remove(wb.active)

    overall = aggregate(url_rows, lambda r: "全部选定路径")
    by_group = aggregate(url_rows, lambda r: r["group"])
    write_sheet(
        wb,
        "Summary",
        SUMMARY_HEADERS,
        [summary_row_values(r) for r in overall + by_group],
        {"范围": 46},
    )
    write_sheet(wb, "URL_Gainers", URL_HEADERS, [url_row_values(r) for r in gainers], {"URL": 60, "结论": 52})
    write_sheet(wb, "URL_Losers", URL_HEADERS, [url_row_values(r) for r in losers], {"URL": 60, "结论": 52})
    write_sheet(
        wb,
        "URL_Impr_Changes",
        URL_HEADERS,
        [url_row_values(r) for r in impr_gainers] + [url_row_values(r) for r in impr_losers],
        {"URL": 60, "结论": 52},
    )
    write_sheet(
        wb,
        "KW_Gainers",
        KW_HEADERS,
        [kw_row_values(r) for r in kw_slice(gainers, False)],
        {"URL": 60, "关键词": 40},
    )
    write_sheet(
        wb,
        "KW_Losers",
        KW_HEADERS,
        [kw_row_values(r) for r in kw_slice(losers, True)],
        {"URL": 60, "关键词": 40},
    )

    notes = [
        ["分析域名 domain_id", domain_id],
        ["本周", f"{cur_start} ~ {cur_end}"],
        ["上周", f"{prev_start} ~ {prev_end}"],
        ["路径前缀", " / ".join(args.paths)],
        ["归组规则", "按最长前缀归组，blogs 与 collections 优先，其余归入根路径，不重复计数"],
        ["URL 级口径", "url_gsc_traffic 的 url_* 字段，先按 (url_id, data_date) 去重后按天求和"],
        ["关键词级口径", "keyword_* 字段，仅 GSC 头部查询，合计小于 URL 总量，只用于归因"],
        ["点击率", "clicks / impressions（加权），单位 %；差值单位为百分点 pp"],
        ["平均排名", "SUM(rank × impressions) / SUM(impressions)，越小越好，差值为负表示排名前进"],
        ["变化分解", "Δclicks = Δ展现×上周CTR + 上周展现×ΔCTR + Δ展现×ΔCTR"],
        ["入榜门槛", f"两周展现量最大值 ≥ {args.min_impr}"],
        ["榜单条数", f"升/降各 {args.top} 条，每个 URL 展开 {args.kw_top} 个关键词"],
    ]
    write_sheet(wb, "Notes", ["项目", "说明"], notes, {"项目": 24, "说明": 90})
    wb.save(args.out)

    # ---------------- 文字结论 ----------------
    lines = [f"# GSC 周对比分析（{cur_start} ~ {cur_end} vs {prev_start} ~ {prev_end}）\n"]
    o = overall[0]
    lines.append(
        f"整体：展现 {int(o['prev_impr']):,} → {int(o['cur_impr']):,}（{o['d_impr']:+,.0f}），"
        f"点击 {int(o['prev_clicks']):,} → {int(o['cur_clicks']):,}（{o['d_clicks']:+,.0f}），"
        f"CTR {o['prev_ctr'] * 100:.2f}% → {o['cur_ctr'] * 100:.2f}%，"
        f"平均排名 {o['prev_rank']:.2f} → {o['cur_rank']:.2f}。"
        f"点击变化中展现效应 {o['impr_effect']:+.0f}、点击率效应 {o['ctr_effect']:+.0f}、"
        f"交互项 {o['interaction']:+.0f}，主因是 {o['driver']}。\n"
    )
    lines.append("## 分路径")
    for g in by_group:
        lines.append(
            f"- `{g['key']}`（{g['urls']} 个 URL）：点击 {int(g['prev_clicks']):,} → {int(g['cur_clicks']):,}"
            f"（{g['d_clicks']:+,.0f}），展现 {g['d_impr']:+,.0f}，CTR {g['d_ctr'] * 100:+.2f}pp，"
            f"平均排名 {g['d_rank']:+.2f}，主因 {g['driver']}。"
        )
    for title, rows_, ascending in (
        ("## 提升最多的 URL", gainers, False),
        ("## 下降最多的 URL", losers, True),
    ):
        lines.append(f"\n{title}")
        for r in rows_[:10]:
            words = sorted(
                kw_by_url.get(r["url_id"], []),
                key=lambda w: w["d_clicks"] if ascending else -w["d_clicks"],
            )[:3]
            kw_text = (
                "；".join(
                    f"{w['keyword']}（点击 {w['d_clicks']:+.0f}，展现 {w['d_impr']:+.0f}，"
                    f"排名 {w['d_rank']:+.2f}，{w['driver']}）"
                    for w in words
                )
                or "无头部关键词数据"
            )
            lines.append(
                f"- {r['url']}\n"
                f"  点击 {int(r['prev_clicks'])} → {int(r['cur_clicks'])}（{r['d_clicks']:+.0f}），"
                f"展现 {int(r['prev_impr'])} → {int(r['cur_impr'])}（{r['d_impr']:+.0f}），"
                f"CTR {r['prev_ctr'] * 100:.2f}% → {r['cur_ctr'] * 100:.2f}%，"
                f"排名 {r['prev_rank']:.2f} → {r['cur_rank']:.2f}。\n"
                f"  归因：{r['explain']}（展现效应 {r['impr_effect']:+.0f} / 点击率效应 "
                f"{r['ctr_effect']:+.0f} / 交互 {r['interaction']:+.0f}）。\n"
                f"  关键词：{kw_text}"
            )
    text = "\n".join(lines)
    if args.markdown:
        with open(args.markdown, "w", encoding="utf-8") as fh:
            fh.write(text + "\n")
    print(text)
    print(f"\nExcel 已输出：{args.out}")


if __name__ == "__main__":
    main()
